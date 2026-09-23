"""
Gera planilhas mensais FICTICIAS no layout de 52 colunas (aba "Tabela").

Todos os nomes, valores, taxas e indicadores sao inventados.
Uso:  python gerar_dados_ficticios.py      -> gera os arquivos em ../exemplos
"""
import random
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SEMENTE = 42  # mesma semente = mesmos dados em toda execucao (reprodutivel)

CABECALHOS = [
    "Data_Base", "ID_Aplicacao", "Fonte", "Centro_Custo", "Instituicao",
    "Modalidade", "Produto", "Indexador", "Perc_CDI_Contratado", "Taxa_Prefixada",
    "Data_Aplicacao", "Data_Vencimento", "Liquidez", "Classificacao", "Saldo_Anterior",
    "Aplicacoes", "Resgates", "Juros_Recebidos", "Rendimento_Bruto", "Saldo_Final",
    "IR_Retido", "IOF_Retido", "IR_Provisionado", "IOF_Provisionado", "Saldo_Liquido",
    "Rentab_Mes", "Rentab_Perc_CDI", "Rentab_Ano", "Rentab_12m", "CDI_Mes",
    "CDI_Ano", "CDI_12m", "IPCA_Mes", "Dias_Uteis_Mes", "Prazo_Dias",
    "Faixa_Vencimento", "Saldo_Liquido_de_IR", "Aliquota_IR", "Isento_IR", "Garantia_FGC",
    "Rating_Emissor", "Emissor", "Custodiante", "Conta_Referencia", "Tipo_Resgate",
    "Carencia_Dias", "Enquadramento", "Observacao", "Area_Responsavel", "Data_Atualizacao",
    "Status", "Chave_Unica",
]
assert len(CABECALHOS) == 52

# Indicadores de mercado FICTICIOS (fracoes). Marco tem "deflacao" para demonstrar o efeito na meta.
MESES = [
    # (ano, mes, data-base registrada na planilha, CDI mes, IPCA mes, dias uteis)
    (2025, 1, date(2025, 1, 31), 0.0095, 0.0030, 22),
    (2025, 2, date(2025, 2, 28), 0.0088, 0.0045, 18),
    (2025, 3, date(2025, 3, 31), 0.0097, -0.0010, 21),
    (2025, 4, date(2025, 4, 26), 0.0092, 0.0035, 20),  # sabado: demonstra ALERTA
    (2025, 5, date(2025, 5, 30), 0.0100, 0.0040, 21),
]

FONTES = ["Recursos Proprios", "Recursos de Terceiros", "Financiamento de Base"]
INSTITUICOES = ["Banco Alfa", "Banco Beta", "Banco Gama", "Cooperativa Delta", "Banco Epsilon"]
# modalidade, indexador, faixa de taxa, isento de IR, saldo informado liquido de IR
MODALIDADES = [
    ("CDB", "CDI", (0.98, 1.05), "N", "N"),
    ("LCA", "CDI", (0.88, 0.95), "S", "N"),
    ("LCI", "CDI", (0.87, 0.94), "S", "N"),
    ("Fundo DI", "CDI", (0.94, 0.99), "N", "S"),
    ("Fundo Aplic. Automatica", "CDI", (0.20, 0.35), "N", "S"),
    ("Tesouro Selic", "CDI", (0.99, 1.01), "N", "N"),
    ("Tesouro IPCA+", "IPCA", (0.0045, 0.0055), "N", "N"),  # spread mensal sobre o IPCA
    ("Compromissada", "CDI", (0.80, 0.90), "N", "N"),
]


def fim_do_mes(ano, mes):
    prox = date(ano + (mes == 12), mes % 12 + 1, 1)
    return date.fromordinal(prox.toordinal() - 1)


def criar_carteira(rng, n=28):
    carteira = []
    for i in range(1, n + 1):
        mod, idx, faixa, isento, liquido = rng.choice(MODALIDADES)
        inst = rng.choice(INSTITUICOES)
        aplic = date(2024, rng.randint(1, 12), rng.randint(1, 28))
        sem_venc = mod.startswith("Fundo") or mod == "Compromissada"
        venc = None if sem_venc else date(rng.randint(2025, 2029), rng.randint(1, 12), rng.randint(1, 28))
        carteira.append(dict(
            id=f"APL-{i:03d}", fonte=rng.choice(FONTES), cc=f"CC-{rng.randint(1, 9):03d}",
            inst=inst, mod=mod, idx=idx, taxa=round(rng.uniform(*faixa), 4),
            isento=isento, liquido=liquido, aplic=aplic, venc=venc,
            saldo=round(rng.uniform(50_000, 2_500_000), 2),
        ))
    return carteira


def faixa_venc(prazo):
    if prazo is None:
        return "Liquidez diaria"
    if prazo <= 90:
        return "Ate 90 dias"
    if prazo <= 360:
        return "91 a 360 dias"
    return "Acima de 360 dias"


def gerar_mes(rng, carteira, ano, mes, data_registrada, cdi, ipca, du, inst_extra=None):
    fim = fim_do_mes(ano, mes)
    linhas = []
    for k, a in enumerate(carteira):
        sa = a["saldo"]
        ap = round(rng.uniform(10_000, 300_000), 2) if rng.random() < 0.15 else 0.0
        rs = round(min(sa, rng.uniform(10_000, 400_000)), 2) if rng.random() < 0.20 else 0.0
        taxa_mes = (ipca + a["taxa"]) if a["idx"] == "IPCA" else cdi * a["taxa"]
        rend = round((sa + ap / 2 - rs / 2) * taxa_mes, 2)
        juros = round(sa * 0.02, 2) if (a["mod"] == "Tesouro IPCA+" and mes in (1, 7)) else 0.0
        aliq = 0.0 if a["isento"] == "S" else 0.15
        ir_ret = round(max(rend, 0) * aliq, 2) if a["liquido"] == "S" else 0.0
        iof_ret = 0.0
        # Fundos com saldo liquido: o IR retido na cota ja sai do saldo final
        sf = round(sa + ap - rs - juros + rend - (ir_ret + iof_ret), 2)
        ir_prov = 0.0 if a["liquido"] == "S" else round(max(rend, 0) * aliq, 2)
        rent = rend / (sa + ap) if (sa + ap) else 0.0
        prazo = (a["venc"] - fim).days if a["venc"] else None
        inst = inst_extra if (inst_extra and k == 0) else a["inst"]
        linhas.append([
            data_registrada, a["id"], a["fonte"], a["cc"], inst,
            a["mod"], f"{a['mod']} {inst}", a["idx"],
            a["taxa"] if a["idx"] == "CDI" else None,
            a["taxa"] if a["idx"] == "IPCA" else None,
            a["aplic"], a["venc"], "Diaria" if a["venc"] is None else "No vencimento",
            "Curto prazo" if (prazo is None or prazo <= 360) else "Longo prazo",
            sa, ap, rs, juros, rend, sf,
            ir_ret, iof_ret, ir_prov, 0.0, round(sf - ir_prov, 2),
            round(rent, 6), round(rent / cdi, 4) if cdi else None, None, None, cdi,
            None, None, ipca, du, prazo,
            faixa_venc(prazo), a["liquido"], aliq, a["isento"],
            "S" if a["mod"] in ("CDB", "LCA", "LCI") else "N",
            rng.choice(["AAA", "AA+", "AA"]), inst, "Custodiante Ficticio",
            f"CT-{k + 1:03d}", "D+0" if a["venc"] is None else "No vencimento",
            0, "Autorizado", "", "Financeiro", fim,
            "Ativa" if sf > 0 else "Encerrada", f"{a['id']}-{ano}{mes:02d}",
        ])
        a["saldo"] = sf
    return linhas


def salvar_planilha(linhas, caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = "Tabela"
    ws.append(CABECALHOS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Arial")
        c.fill = PatternFill("solid", fgColor="44546A")
    for ln in linhas:
        ws.append(ln)
    for col in ("A", "K", "L", "AX"):
        for cel in ws[col][1:]:
            cel.number_format = "dd/mm/yyyy"
    for col in ("O", "P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y"):
        for cel in ws[col][1:]:
            cel.number_format = "#,##0.00"
    ws.freeze_panes = "A2"
    wb.save(caminho)


def gerar_todos(destino=None):
    """Gera todos os meses; se 'destino' for informado, salva um .xlsx por mes."""
    rng = random.Random(SEMENTE)
    carteira = criar_carteira(rng)
    resultado = {}
    for ano, mes, dreg, cdi, ipca, du in MESES:
        extra = "Banco Zeta" if (ano, mes) == (2025, 5) else None  # instituicao nao cadastrada -> ERRO
        linhas = gerar_mes(rng, carteira, ano, mes, dreg, cdi, ipca, du, extra)
        sufixo = "_com_erro" if extra else ("_data_fim_de_semana" if dreg.weekday() >= 5 else "")
        if destino:
            salvar_planilha(linhas, Path(destino) / f"planilha_mensal_{ano}-{mes:02d}{sufixo}.xlsx")
        resultado[(ano, mes)] = linhas
    return resultado


if __name__ == "__main__":
    pasta = Path(__file__).resolve().parent.parent / "exemplos"
    pasta.mkdir(exist_ok=True)
    gerar_todos(pasta)
    print("Planilhas ficticias geradas em", pasta)
