"""
Constroi excel/base_relatorio_exemplo.xlsx com dados FICTICIOS (jan-mar/2025 pre-carregados).

A base e o "motor" do relatorio: tabela historica (tblBase), cadastros, indicadores,
log, conferencias OK/ALERTA/ERRO e resumo executivo, tudo por formulas.
Depois de gerar: abra no Excel, salve como .xlsm e importe os modulos da pasta /vba.
"""
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.table import Table, TableStyleInfo

from gerar_dados_ficticios import CABECALHOS, FONTES, INSTITUICOES, MESES, MODALIDADES, fim_do_mes, gerar_todos

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "excel" / "base_relatorio_exemplo.xlsx"
MESES_PRE = [(2025, 1), (2025, 2), (2025, 3)]
N = 50000  # extensao dos intervalos nomeados da base

# Paleta NEUTRA e generica (nao institucional)
COR_PRI, COR_SEC, COR_CLARA = "1F3A5F", "44546A", "EEF2F7"
F = "Arial"
fonte_tit = Font(name=F, size=16, bold=True, color=COR_PRI)
fonte_sub = Font(name=F, size=11, bold=True, color="FFFFFF")
fonte_n = Font(name=F, size=10)
fonte_b = Font(name=F, size=10, bold=True)
fonte_input = Font(name=F, size=10, color="0000FF")
fill_sub = PatternFill("solid", fgColor=COR_SEC)
fill_cl = PatternFill("solid", fgColor=COR_CLARA)
fill_in = PatternFill("solid", fgColor="FFF2CC")
fina = Side(style="thin", color="BFBFBF")
borda = Border(top=fina, bottom=fina, left=fina, right=fina)
MOEDA = '"R$" #,##0.00;-"R$" #,##0.00;"-"'
PCT = "0.00%"
MESES_PT = '"janeiro","fevereiro","março","abril","maio","junho","julho","agosto","setembro","outubro","novembro","dezembro"'


def nome(wb, n, ref):
    wb.defined_names[n] = DefinedName(n, attr_text=ref)


def secao(ws, cel, texto, ate):
    ws[cel] = texto
    ws.merge_cells(f"{cel}:{ate}")
    ws[cel].font = fonte_sub
    ws[cel].fill = fill_sub
    ws[cel].alignment = Alignment(vertical="center", indent=1)


def cab(ws, linha, col_ini, titulos):
    for i, t in enumerate(titulos):
        c = ws.cell(row=linha, column=col_ini + i, value=t)
        c.font = fonte_b
        c.fill = fill_cl
        c.border = borda
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def fmt(ws, rng, formato=None, negrito=False):
    for row in ws[rng]:
        for c in row:
            c.font = fonte_b if negrito else fonte_n
            c.border = borda
            if formato:
                c.number_format = formato


def construir():
    dados = gerar_todos(None)
    wb = Workbook()

    # ---------------- LEIAME ----------------
    ws = wb.active
    ws.title = "LEIAME"
    ws.column_dimensions["A"].width = 110
    linhas = [
        ("Gerador de Relatório Executivo de Carteira de Renda Fixa — ARQUIVO DE EXEMPLO", fonte_tit),
        ("Todos os dados deste arquivo são FICTÍCIOS e servem apenas para demonstração.", fonte_b),
        ("", None),
        ("Como usar:", fonte_b),
        ("1. Salve este arquivo como 'Pasta de Trabalho Habilitada para Macro (*.xlsm)'.", None),
        ("2. Abra o editor VBA (Alt+F11), selecione ESTE arquivo no painel e importe os 3 módulos da pasta /vba.", None),
        ("3. Execute a macro CriarBotoes (Alt+F8) uma única vez para criar os botões na aba Resumo.", None),
        ("4. Clique em 'Importar mês' e escolha exemplos/planilha_mensal_2025-04_data_fim_de_semana.xlsx.", None),
        ("5. Veja a aba Conferencias: deve aparecer ALERTA de data-base em fim de semana (não bloqueia o PDF).", None),
        ("6. Importe exemplos/planilha_mensal_2025-05_com_erro.xlsx: deve aparecer ERRO (instituição não cadastrada).", None),
        ("7. Clique em 'Gerar PDF'. Com ERRO, a macro só permite gerar a versão PRELIMINAR.", None),
        ("", None),
        ("Células com fundo amarelo e fonte azul são entradas manuais. As demais são fórmulas: não sobrescreva.", fonte_b),
        ("A macro NÃO salva a base automaticamente: revise e salve manualmente após cada importação.", None),
    ]
    for i, (t, f) in enumerate(linhas, 1):
        ws.cell(row=i, column=1, value=t).font = f or fonte_n

    # ---------------- BASE ----------------
    wsb = wb.create_sheet("Base")
    colunas = CABECALHOS + ["Mov_Diferenca", "Saldo_Medio"]
    wsb.append(colunas)
    r = 2
    for chave in MESES_PRE:
        fim = fim_do_mes(*chave)
        for ln in dados[chave]:
            ln = list(ln)
            ln[0] = fim  # data-base normalizada para o ultimo dia do mes
            wsb.append(ln + [
                f'=ROUND(T{r}-(O{r}+P{r}-Q{r}-R{r}+S{r}-IF(AK{r}="S",U{r}+V{r},0)),2)',
                f"=(O{r}+T{r})/2",
            ])
            r += 1
    ult = r - 1
    t = Table(displayName="tblBase", ref=f"A1:BB{ult}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    wsb.add_table(t)
    for col in ("A", "K", "L", "AX"):
        for c in wsb[col][1:]:
            c.number_format = "dd/mm/yyyy"
    for col in ("O", "P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "BA", "BB"):
        for c in wsb[col][1:]:
            c.number_format = "#,##0.00"
    wsb.freeze_panes = "C2"
    wsb["BA1"].comment = Comment("Coluna auxiliar: saldo final menos a movimentação calculada. "
                                 "Fundos com saldo líquido de IR (coluna AK = S) descontam o IR/IOF retido.", "Modelo")

    col = dict(Data="A", ID="B", Fonte="C", Inst="E", Mod="F", SA="O", Ap="P", Res="Q", Jur="R",
               Rend="S", SF="T", IRR="U", IOFR="V", IRP="W", Mov="BA", SMed="BB")
    for k, letra in col.items():
        nome(wb, f"col_{k}", f"Base!${letra}$2:${letra}${N}")

    # ---------------- CONFIG ----------------
    wc = wb.create_sheet("Config")
    wc.column_dimensions["A"].width = 44
    wc.column_dimensions["B"].width = 22
    wc.column_dimensions["C"].width = 70
    wc["A1"] = "Parâmetros do relatório"
    wc["A1"].font = fonte_tit
    params = [
        ("Mês de referência", date(2025, 3, 31), "dd/mm/yyyy", "Atualizado pela macro ao importar. Pode ser alterado para rever meses anteriores."),
        ("Limite de concentração por instituição", 0.30, "0%", "Parâmetro FICTÍCIO. Substitua pelo limite da sua política."),
        ("Meta de rentabilidade (% do CDI)", 1.00, "0%", "Parâmetro FICTÍCIO."),
        ("Tolerância de continuidade do saldo (R$)", 1.00, "#,##0.00", "Diferença aceitável entre saldo final anterior e saldo inicial atual."),
        ("Pasta do PDF (vazio = pasta deste arquivo)", "", "@", "Caminho local. Se o arquivo estiver em nuvem, a macro pergunta a pasta."),
        ("Status do relatório", "FINAL", "@", "Preenchido pela macro (FINAL ou PRELIMINAR). Não editar."),
    ]
    for i, (lab, val, f_, obs) in enumerate(params, 3):
        wc.cell(row=i, column=1, value=lab).font = fonte_b
        c = wc.cell(row=i, column=2, value=val)
        c.number_format = f_
        c.font, c.fill, c.border = fonte_input, fill_in, borda
        wc.cell(row=i, column=3, value=obs).font = Font(name=F, size=9, italic=True, color="595959")
    wc["A9"] = "Mês anterior (calculado)"
    wc["A9"].font = fonte_b
    wc["B9"] = "=EOMONTH(B3,-1)"
    wc["B9"].number_format = "dd/mm/yyyy"
    for n_, ref in [("MesRef", "$B$3"), ("LimiteConc", "$B$4"), ("MetaCDI", "$B$5"),
                    ("TolContinuidade", "$B$6"), ("PastaPDF", "$B$7"), ("StatusRelatorio", "$B$8"), ("MesAnt", "$B$9")]:
        nome(wb, n_, f"Config!{ref}")

    # ---------------- CADASTROS ----------------
    wk = wb.create_sheet("Cadastros")
    listas = [("A", "Instituicao", INSTITUICOES, "tblInstituicoes", "cad_Inst"),
              ("C", "Modalidade", [m[0] for m in MODALIDADES], "tblModalidades", "cad_Mod"),
              ("E", "Fonte", FONTES, "tblFontes", "cad_Fontes")]
    for letra, titulo, itens, tbl, nm in listas:
        wk[f"{letra}1"] = titulo
        for i, it in enumerate(itens, 2):
            wk[f"{letra}{i}"] = it
        tb = Table(displayName=tbl, ref=f"{letra}1:{letra}{len(itens) + 1}")
        tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        wk.add_table(tb)
        wk.column_dimensions[letra].width = 28
        fim_nm = len(itens) + 1 if nm == "cad_Fontes" else 200
        nome(wb, nm, f"Cadastros!${letra}$2:${letra}${fim_nm}")
    wk["G1"] = "Mes"
    for i, (a, m) in enumerate(MESES_PRE, 2):
        wk[f"G{i}"] = fim_do_mes(a, m)
        wk[f"G{i}"].number_format = "mm/yyyy"
    tb = Table(displayName="tblMeses", ref=f"G1:G{len(MESES_PRE) + 1}")
    tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    wk.add_table(tb)
    wk.column_dimensions["G"].width = 14
    nome(wb, "cad_Meses", "Cadastros!$G$2:$G$500")
    wk["I1"] = "Novas instituições ou modalidades: acrescente uma linha na tabela correspondente (a conferência acusa ERRO até lá)."
    wk["I1"].font = Font(name=F, size=9, italic=True)

    # ---------------- INDICADORES ----------------
    wi = wb.create_sheet("Indicadores")
    wi.append(["Mes", "CDI_Mes", "IPCA_Mes", "Data_Base_Original"])
    for a, m, dreg, cdi, ipca, du in MESES:
        if (a, m) in MESES_PRE:
            wi.append([fim_do_mes(a, m), cdi, ipca, dreg])
    tb = Table(displayName="tblIndicadores", ref=f"A1:D{len(MESES_PRE) + 1}")
    tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    wi.add_table(tb)
    for row in wi.iter_rows(min_row=2):
        row[0].number_format = "mm/yyyy"
        row[1].number_format = row[2].number_format = "0.00%"
        row[3].number_format = "dd/mm/yyyy"
    for c_, w in zip("ABCD", (12, 12, 12, 20)):
        wi.column_dimensions[c_].width = w
    wi["F1"] = "Valores FICTÍCIOS, gravados pela macro a partir das colunas CDI_Mes e IPCA_Mes da planilha mensal."
    wi["F1"].font = Font(name=F, size=9, italic=True)
    for n_, l in [("ind_Mes", "A"), ("ind_CDI", "B"), ("ind_IPCA", "C"), ("ind_Orig", "D")]:
        nome(wb, n_, f"Indicadores!${l}$2:${l}$500")

    # ---------------- LOG ----------------
    wl = wb.create_sheet("Log")
    wl.append(["Data_Hora", "Usuario", "Arquivo", "Mes_Ref", "Linhas", "Data_Base_Original", "Observacao"])
    wl.append(["", "script", "construir_base.py", fim_do_mes(2025, 3), ult - 1, "", "Carga inicial de exemplo (jan-mar/2025, dados fictícios)"])
    tb = Table(displayName="tblLog", ref="A1:G2")
    tb.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    wl.add_table(tb)
    wl["D2"].number_format = "mm/yyyy"
    for c_, w in zip("ABCDEFG", (18, 14, 44, 10, 8, 18, 60)):
        wl.column_dimensions[c_].width = w

    # ---------------- CONFERENCIAS ----------------
    wf = wb.create_sheet("Conferencias", 1)
    for c_, w in zip("ABCDE", (5, 48, 16, 44, 12)):
        wf.column_dimensions[c_].width = w
    wf["A1"] = "Conferências de qualidade dos dados"
    wf["A1"].font = fonte_tit
    wf["B3"], wf["B4"] = "Quantidade de ERROS (bloqueiam o PDF final)", "Quantidade de ALERTAS (não bloqueiam)"
    wf["C3"], wf["C4"] = '=COUNTIF(E8:E30,"ERRO")', '=COUNTIF(E8:E30,"ALERTA")'
    for c_ in ("B3", "B4", "C3", "C4"):
        wf[c_].font = fonte_b
    nome(wb, "QtdErros", "Conferencias!$C$3")
    nome(wb, "QtdAlertas", "Conferencias!$C$4")
    cab(wf, 7, 1, ["#", "Conferência", "Resultado", "Critério", "Status"])
    D, M = "col_Data", "MesRef"
    confs = [
        ("Mês de referência cadastrado", f"=COUNTIF(cad_Meses,{M})", "Deve ser ≥ 1", "=IF(C{r}>=1,\"OK\",\"ERRO\")"),
        ("Linhas do mês na base", f"=COUNTIF({D},{M})", "Deve ser > 0", "=IF(C{r}>0,\"OK\",\"ERRO\")"),
        ("Linhas sem instituição ou modalidade", f'=SUMPRODUCT(({D}={M})*(col_Inst=""))+SUMPRODUCT(({D}={M})*(col_Mod=""))', "Deve ser 0", "=IF(C{r}=0,\"OK\",\"ERRO\")"),
        ("Soma por fonte × total (diferença R$)", f"=ROUND(SUMIFS(col_SF,{D},{M})-SUMPRODUCT(SUMIFS(col_SF,{D},{M},col_Fonte,cad_Fontes)),2)", "Deve ser 0,00", "=IF(ABS(C{r})<0.01,\"OK\",\"ERRO\")"),
        ("Instituições não cadastradas (linhas)", f"=SUMPRODUCT(({D}={M})*(COUNTIF(cad_Inst,col_Inst)=0))", "Deve ser 0", "=IF(C{r}=0,\"OK\",\"ERRO\")"),
        ("Modalidades não cadastradas (linhas)", f"=SUMPRODUCT(({D}={M})*(COUNTIF(cad_Mod,col_Mod)=0))", "Deve ser 0", "=IF(C{r}=0,\"OK\",\"ERRO\")"),
        ("Linhas com movimentação que não fecha", f"=SUMPRODUCT(({D}={M})*(ABS(col_Mov)>0.01))", "Deve ser 0 (tolerância R$ 0,01)", "=IF(C{r}=0,\"OK\",\"ERRO\")"),
        ("Continuidade do saldo (diferença R$)", f'=IF(COUNTIF({D},MesAnt)=0,"sem mês ant.",ROUND(SUMIFS(col_SA,{D},{M})-SUMIFS(col_SF,{D},MesAnt),2))', "Até a tolerância da aba Config", "=IF(ISNUMBER(C{r}),IF(ABS(C{r})<=TolContinuidade,\"OK\",\"ALERTA\"),\"ALERTA\")"),
        ("Data-base original em dia útil", f'=IFERROR(INDEX(ind_Orig,MATCH({M},ind_Mes,0)),"não localizada")', "Segunda a sexta (feriados não considerados)", "=IF(ISNUMBER(C{r}),IF(WEEKDAY(C{r},2)<=5,\"OK\",\"ALERTA\"),\"ALERTA\")"),
        ("Aplicações com rendimento negativo", f"=SUMPRODUCT(({D}={M})*(col_Rend<0))", "Esperado 0", "=IF(C{r}=0,\"OK\",\"ALERTA\")"),
        ("Mês anterior disponível na base", f"=COUNTIF({D},MesAnt)", "Deve ser > 0", "=IF(C{r}>0,\"OK\",\"ALERTA\")"),
        ("CDI do mês disponível", f'=IFERROR(INDEX(ind_CDI,MATCH({M},ind_Mes,0)),"não localizado")', "Deve estar preenchido", "=IF(AND(ISNUMBER(C{r}),C{r}<>0),\"OK\",\"ERRO\")"),
    ]
    for i, (desc, f_res, crit, f_st) in enumerate(confs, 1):
        r = 7 + i
        wf.cell(row=r, column=1, value=i)
        wf.cell(row=r, column=2, value=desc)
        wf.cell(row=r, column=3, value=f_res)
        wf.cell(row=r, column=4, value=crit)
        wf.cell(row=r, column=5, value=f_st.format(r=r))
    fmt(wf, "A8:E19")
    wf["C11"].number_format = wf["C15"].number_format = "#,##0.00"
    wf["C16"].number_format = "dd/mm/yyyy"
    wf["C19"].number_format = "0.00%"
    for c in wf["E"][7:19]:
        c.alignment = Alignment(horizontal="center")
        c.font = fonte_b
    verm, amar, verd = (PatternFill("solid", fgColor=x) for x in ("F8CBAD", "FFE699", "C6EFCE"))
    wf.conditional_formatting.add("E8:E19", CellIsRule(operator="equal", formula=['"ERRO"'], fill=verm))
    wf.conditional_formatting.add("E8:E19", CellIsRule(operator="equal", formula=['"ALERTA"'], fill=amar))
    wf.conditional_formatting.add("E8:E19", CellIsRule(operator="equal", formula=['"OK"'], fill=verd))

    # ---------------- RESUMO ----------------
    wr = wb.create_sheet("Resumo", 1)
    for c_, w in zip("ABCDEFG", (40, 18, 18, 18, 14, 14, 16)):
        wr.column_dimensions[c_].width = w
    wr["A1"] = "Relatório Executivo de Investimentos"
    wr["A1"].font = fonte_tit
    wr["A2"] = f'="Mês de referência: "&CHOOSE(MONTH(MesRef),{MESES_PT})&"/"&YEAR(MesRef)&"   |   Dados fictícios de demonstração"'
    wr["A2"].font = fonte_b
    wr["A3"] = '=IF(StatusRelatorio="PRELIMINAR","VERSÃO PRELIMINAR — há conferências com ERRO","")'
    wr["A3"].font = Font(name=F, size=11, bold=True, color="C00000")

    secao(wr, "A5", "1. Resultado do mês", "D5")
    cab(wr, 6, 1, ["Indicador", "Mês de referência", "Mês anterior", "Variação"])
    def s(c, mes):
        return f"SUMIFS(col_{c},col_Data,{mes})"
    kp = [
        ("Saldo final total", "SF", MOEDA), ("Aplicações", "Ap", MOEDA), ("Resgates", "Res", MOEDA),
        ("Juros recebidos", "Jur", MOEDA), ("Rendimento bruto", "Rend", MOEDA),
        ("IR/IOF retido na cota", None, MOEDA), ("IR provisionado", "IRP", MOEDA), ("Saldo médio (média simples)", "SMed", MOEDA),
    ]
    for i, (lab, c, f_) in enumerate(kp, 7):
        wr.cell(row=i, column=1, value=lab)
        for j, mes in ((2, "MesRef"), (3, "MesAnt")):
            val = f"={s('IRR', mes)}+{s('IOFR', mes)}" if c is None else f"={s(c, mes)}"
            wr.cell(row=i, column=j, value=val).number_format = f_
        wr.cell(row=i, column=4, value=f"=B{i}-C{i}").number_format = f_
    # linhas 15-19: rentabilidade
    wr["A15"], wr["B15"], wr["C15"] = "Rentabilidade do mês", "=IFERROR(B11/B14,0)", "=IFERROR(C11/C14,0)"
    wr["A16"] = "CDI do mês"
    wr["B16"] = "=IFERROR(INDEX(ind_CDI,MATCH(MesRef,ind_Mes,0)),0)"
    wr["C16"] = "=IFERROR(INDEX(ind_CDI,MATCH(MesAnt,ind_Mes,0)),0)"
    wr["A17"], wr["B17"], wr["C17"] = "Rentabilidade em % do CDI", "=IFERROR(B15/B16,0)", "=IFERROR(C15/C16,0)"
    wr["A18"], wr["B18"], wr["C18"] = "Meta (% do CDI)", "=MetaCDI", "=MetaCDI"
    wr["A19"] = "Situação da meta"
    wr["B19"] = '=IF(B17>=B18,"Atingida","Não atingida")'
    wr["C19"] = '=IF(C16=0,"-",IF(C17>=C18,"Atingida","Não atingida"))'
    for rr in (15, 16, 17, 18):
        wr[f"D{rr}"] = f"=B{rr}-C{rr}"
        for cc in "BCD":
            wr[f"{cc}{rr}"].number_format = PCT
    fmt(wr, "A7:D19")
    for rr in range(7, 20):
        for cc in "BCD":
            if rr in (15, 16, 17, 18):
                wr[f"{cc}{rr}"].number_format = PCT
            elif rr != 19:
                wr[f"{cc}{rr}"].number_format = MOEDA

    secao(wr, "A21", "2. Variação do rendimento: efeito volume × efeito taxa", "D21")
    dec = [
        ("Variação do rendimento bruto", "=B11-C11"),
        ("Efeito volume = (saldo médio atual − anterior) × taxa anterior", "=IFERROR((B14-C14)*C15,0)"),
        ("Efeito taxa = variação − efeito volume", "=B22-B23"),
    ]
    for i, (lab, f_) in enumerate(dec, 22):
        wr.cell(row=i, column=1, value=lab)
        wr.cell(row=i, column=2, value=f_).number_format = MOEDA
    fmt(wr, "A22:B24", MOEDA)
    for rr in range(22, 25):
        wr[f"A{rr}"].number_format = "General"

    secao(wr, "A26", "3. Distribuição por fonte de recursos", "D26")
    cab(wr, 27, 1, ["Fonte", "Saldo final", "Participação", "Rendimento bruto"])
    for i in range(3):
        rr = 28 + i
        wr[f"A{rr}"] = f"=Cadastros!E{2 + i}"
        wr[f"B{rr}"] = f"=SUMIFS(col_SF,col_Data,MesRef,col_Fonte,A{rr})"
        wr[f"C{rr}"] = f"=IFERROR(B{rr}/$B$7,0)"
        wr[f"D{rr}"] = f"=SUMIFS(col_Rend,col_Data,MesRef,col_Fonte,A{rr})"
    wr["A31"], wr["B31"], wr["C31"], wr["D31"] = "Total", "=SUM(B28:B30)", "=SUM(C28:C30)", "=SUM(D28:D30)"
    fmt(wr, "A28:D30")
    fmt(wr, "A31:D31", negrito=True)
    for rr in range(28, 32):
        wr[f"B{rr}"].number_format = wr[f"D{rr}"].number_format = MOEDA
        wr[f"C{rr}"].number_format = PCT

    secao(wr, "A33", "4. Concentração por instituição", "F33")
    cab(wr, 34, 1, ["Instituição", "Saldo final", "Participação", "Limite", "Margem", "Situação"])
    for i in range(10):
        rr = 35 + i
        wr[f"A{rr}"] = f'=IF(Cadastros!A{2 + i}="","",Cadastros!A{2 + i})'
        wr[f"B{rr}"] = f'=IF(A{rr}="","",SUMIFS(col_SF,col_Data,MesRef,col_Inst,A{rr}))'
        wr[f"C{rr}"] = f'=IF(A{rr}="","",IFERROR(B{rr}/$B$7,0))'
        wr[f"D{rr}"] = f'=IF(A{rr}="","",LimiteConc)'
        wr[f"E{rr}"] = f'=IF(A{rr}="","",D{rr}-C{rr})'
        wr[f"F{rr}"] = f'=IF(A{rr}="","",IF(C{rr}>D{rr},"Atenção","Conforme"))'
    fmt(wr, "A35:F44")
    for rr in range(35, 45):
        wr[f"B{rr}"].number_format = MOEDA
        for cc in "CDE":
            wr[f"{cc}{rr}"].number_format = PCT
    wr.conditional_formatting.add("F35:F44", CellIsRule(operator="equal", formula=['"Atenção"'], fill=amar))
    wr["A45"] = "“Atenção” indica participação acima do limite parametrizado; não equivale a não conformidade confirmada."
    wr["A45"].font = Font(name=F, size=8, italic=True)

    secao(wr, "A47", "5. Distribuição por modalidade", "C47")
    cab(wr, 48, 1, ["Modalidade", "Saldo final", "Participação"])
    for i in range(10):
        rr = 49 + i
        wr[f"A{rr}"] = f'=IF(Cadastros!C{2 + i}="","",Cadastros!C{2 + i})'
        wr[f"B{rr}"] = f'=IF(A{rr}="","",SUMIFS(col_SF,col_Data,MesRef,col_Mod,A{rr}))'
        wr[f"C{rr}"] = f'=IF(A{rr}="","",IFERROR(B{rr}/$B$7,0))'
    fmt(wr, "A49:C58")
    for rr in range(49, 59):
        wr[f"B{rr}"].number_format = MOEDA
        wr[f"C{rr}"].number_format = PCT

    secao(wr, "A60", "6. Pontos de atenção (gerados por fórmula)", "F60")
    pontos = [
        '=IF(B17>=B18,"Meta atingida: ","Meta não atingida: ")&FIXED(B17*100,1)&"% do CDI (meta de "&FIXED(B18*100,0)&"%)."',
        '="Variação do rendimento: R$ "&FIXED(B22,2)&" (efeito volume R$ "&FIXED(B23,2)&"; efeito taxa R$ "&FIXED(B24,2)&")."',
        '="Instituições acima do limite de concentração: "&COUNTIF(F35:F44,"Atenção")&"."',
        '="Conferências: "&QtdErros&" erro(s) e "&QtdAlertas&" alerta(s). Detalhes na aba Conferencias."',
    ]
    for i, p in enumerate(pontos, 61):
        wr[f"A{i}"] = p
        wr[f"A{i}"].font = fonte_n
        wr.merge_cells(f"A{i}:F{i}")

    secao(wr, "A66", "7. Comentário da analista (preenchimento manual)", "F66")
    wr["A67"] = "Digite aqui a análise do mês antes de gerar o PDF."
    wr.merge_cells("A67:F71")
    wr["A67"].font, wr["A67"].fill = fonte_input, fill_in
    wr["A67"].alignment = Alignment(wrap_text=True, vertical="top")
    wr["A73"] = "Relatório gerado a partir de dados fictícios. Validação humana obrigatória antes de qualquer uso formal."
    wr["A73"].font = Font(name=F, size=8, italic=True, color="595959")

    for w_ in (wr, wf):
        w_.page_setup.paperSize = w_.PAPERSIZE_A4
        w_.page_setup.orientation = "portrait"
        w_.page_setup.fitToWidth = 1
        w_.page_setup.fitToHeight = 0
        w_.sheet_properties.pageSetUpPr.fitToPage = True
        w_.oddFooter.center.text = "Página &P de &N"
        w_.sheet_view.showGridLines = False
    wr.print_area = "A1:G73"
    wf.print_area = "A1:E19"

    wb.calculation.fullCalcOnLoad = True  # Excel recalcula tudo ao abrir
    wb.properties.creator = "Modelo de exemplo"
    wb.properties.company = ""
    SAIDA.parent.mkdir(exist_ok=True)
    wb.save(SAIDA)
    print("Base gerada:", SAIDA)


if __name__ == "__main__":
    construir()
